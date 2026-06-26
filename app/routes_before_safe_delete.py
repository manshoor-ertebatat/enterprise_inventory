
from flask import Blueprint, render_template, request, redirect, session, flash
from app.models import db, Product, Movement, User, ActivityLog, ProductSerial
from openpyxl import Workbook
from flask import send_file
from sqlalchemy import func
import io
import os
import jdatetime

bp = Blueprint("main", __name__)

@bp.app_template_filter("jalali")
def jalali_date(value):

    if not value:
        return ""

    return jdatetime.datetime.fromgregorian(
        datetime=value
    ).strftime("%Y/%m/%d %H:%M")

def log_activity(username, action, description):

    try:

        log = ActivityLog(
            username=username,
            action=action,
            description=description
        )

        db.session.add(log)
        db.session.commit()

    except Exception:
        db.session.rollback()

def shamsi_date(value):

    if not value:
        return ""

    return jdatetime.datetime.fromgregorian(
        datetime=value
    ).strftime("%Y/%m/%d %H:%M")

bp.add_app_template_filter(
    shamsi_date,
    "shamsi"
)

@bp.route("/")
def login():
    return render_template("login.html")

@bp.route("/login", methods=["POST"])
def do_login():

    username = request.form["username"]
    password = request.form["password"]

    user = User.query.filter_by(
        username=username,
        password=password
    ).first()

    if user:

        session["user"] = user.username
        session["role"] = user.role

        log_activity(
            user.username,
            "LOGIN",
            "ورود به سیستم"
        )

        return redirect("/dashboard")

    return "نام کاربری یا رمز عبور اشتباه است"

@bp.route("/dashboard")
def dashboard():

    if "user" not in session:
        return redirect("/")

    search = request.args.get("search")

    all_products_query = Product.query.order_by(
        Product.id.desc()
    )

    if search:
        all_products_query = all_products_query.filter(
            Product.name.contains(search)
        )

    all_products = all_products_query.all()

    # فقط 10 کالا در داشبورد نمایش داده می‌شود
    products = all_products[:10]

    movements = Movement.query.order_by(
        Movement.id.desc()
    ).limit(10).all()

    total_products = len(all_products)

    low_stock = len([
        p for p in products
        if p.qty > 0 and p.qty < 5
    ])

    out_of_stock = len([
        p for p in products
        if p.qty == 0
    ])

    serial_counts = {}

    for p in products:

        serial_counts[p.id] = ProductSerial.query.filter_by(
            product_id=p.id,
            status="IN_STOCK"
        ).count()

    products_dict = {}

    for product in Product.query.all():
        products_dict[product.id] = product.name

    return render_template(
        "dashboard.html",
        products=products,
        movements=movements,
        total_products=total_products,
        low_stock=low_stock,
        out_of_stock=out_of_stock,
        serial_counts=serial_counts,
        products_dict=products_dict,
    )

@bp.route("/products")
def products_page():

    if "user" not in session:
        return redirect("/")

    search = request.args.get("search", "").strip()
    page = request.args.get("page", 1, type=int)

    query = Product.query.order_by(Product.id.desc())

    if search:
        query = query.filter(Product.name.contains(search))

    pagination = query.paginate(
        page=page,
        per_page=20,
        error_out=False
    )

    serial_counts = {}

    for product in pagination.items:
        serial_counts[product.id] = ProductSerial.query.filter_by(
            product_id=product.id,
            status="IN_STOCK"
        ).count()

    return render_template(
        "products.html",
        products=pagination.items,
        pagination=pagination,
        search=search,
        serial_counts=serial_counts
    )

@bp.route("/add", methods=["POST"])
def add():

    if "user" not in session:
        return redirect("/")

    if session.get("role") == "viewer":
        flash("شما اجازه ثبت کالا ندارید.", "danger")
        return redirect("/dashboard")

    name = request.form.get("name", "").strip()
    qty_text = request.form.get("qty", "0").strip()
    has_serial = 1 if request.form.get("has_serial") else 0

    if not name:
        flash("نام کالا وارد نشده است.", "danger")
        return redirect("/dashboard")

    try:
        qty = int(qty_text)
    except ValueError:
        flash("موجودی اولیه باید عدد باشد.", "danger")
        return redirect("/dashboard")

    if qty < 0:
        flash("موجودی اولیه نمی‌تواند منفی باشد.", "danger")
        return redirect("/dashboard")

    normalized_name = (
        name
        .replace("ي", "ی")
        .replace("ك", "ک")
        .replace("\u200c", " ")
        .strip()
        .lower()
    )

    for item in Product.query.all():
        existing_name = (
            (item.name or "")
            .replace("ي", "ی")
            .replace("ك", "ک")
            .replace("\u200c", " ")
            .strip()
            .lower()
        )

        if existing_name == normalized_name:
            flash(f"کالایی با نام «{name}» قبلاً ثبت شده است.", "warning")
            return redirect("/dashboard")

    p = Product(
        name=name,
        qty=qty,
        has_serial=has_serial
    )

    db.session.add(p)
    db.session.commit()

    log_activity(
        session["user"],
        "ADD_PRODUCT",
        f"افزودن کالا: {name} | موجودی اولیه: {qty}"
    )

    flash(f"کالای «{name}» با موفقیت ثبت شد.", "success")
    return redirect("/dashboard")

@bp.route("/move", methods=["POST"])
def move():

    product_id = int(request.form["product_id"])
    qty = int(request.form["qty"])
    mtype = request.form["type"]
    receiver_name = request.form.get(
        "receiver_name",
        ""
    )

    customer_name = request.form.get(
        "customer_name",
        ""
    )

    description = request.form.get(
        "description",
        ""
    )

    product = Product.query.get(product_id)

    if mtype == "OUT" and product.qty < qty:
        flash("موجودی کالا برای این خروج کافی نیست.", "danger")
        return redirect("/dashboard")

    if mtype == "OUT" and product.has_serial:
        available_serials = ProductSerial.query.filter_by(
            product_id=product.id,
            status="IN_STOCK"
        ).count()

        if available_serials < qty:
            flash(
                f"تعداد سریال‌های موجود کافی نیست. موجود: {available_serials} - درخواست خروج: {qty}",
                "danger"
            )
            return redirect("/dashboard")

    if mtype == "IN":
        product.qty += qty
    else:
        product.qty -= qty

    m = Movement(
        product_id=product_id,
        qty=qty,
        type=mtype,
        receiver_name=receiver_name,
        customer_name=customer_name,
        description=description,
        created_by=session["user"]
    )
    db.session.add(m)
    db.session.commit()

    if mtype == "IN":

        log_activity(
            session["user"],
            "STOCK_IN",
            f"ورود {qty} عدد از کالا {product.name}"
        )

    else:

        log_activity(
            session["user"],
            "STOCK_OUT",
            f"خروج {qty} عدد از کالا {product.name}"
        )

    action_text = "ورود به انبار" if mtype == "IN" else "خروج از انبار"
    flash(
        f"{action_text} برای کالای «{product.name}» با تعداد {qty} با موفقیت ثبت شد.",
        "success"
    )

    return redirect("/dashboard")
@bp.route("/logout")
def logout():

    if "user" in session:

        log_activity(
            session["user"],
            "LOGOUT",
            "خروج از سیستم"
        )

    session.clear()

    return redirect("/")

@bp.route("/delete/<int:id>")
def delete_product(id):

    if "user" not in session:
        return redirect("/")

    if session.get("role") == "viewer":
        return "Access Denied"

    product = Product.query.get(id)

    if product:

        movement_count = Movement.query.filter_by(
            product_id=product.id
        ).count()

        serial_count = ProductSerial.query.filter_by(
            product_id=product.id
        ).count()

        if movement_count > 0 or serial_count > 0:
            return (
                "این کالا دارای سابقه ورود، خروج یا سریال است و قابل حذف نیست. "
                "در صورت عدم نیاز، موجودی آن را صفر کنید."
            )

        log_activity(
            session["user"],
            "DELETE_PRODUCT",
            f"حذف کالا: {product.name}"
        )

        db.session.delete(product)
        db.session.commit()

    return redirect("/dashboard")

@bp.route("/edit/<int:id>", methods=["GET","POST"])
def edit_product(id):

    if "user" not in session:
        return redirect("/")

    if session.get("role") == "viewer":
        return "Access Denied"

    product = Product.query.get(id)

    if request.method == "POST":

        old_name = product.name
        old_qty = product.qty

        product.name = request.form["name"]
        new_qty = int(request.form["qty"])

        new_has_serial = 1 if request.form.get("has_serial") == "1" else 0

        if new_has_serial:
            in_stock_serial_count = ProductSerial.query.filter_by(
                product_id=product.id,
                status="IN_STOCK"
            ).count()

            if new_qty < in_stock_serial_count:
                return (
                    f"موجودی جدید نمی‌تواند کمتر از تعداد سریال‌های موجود باشد. "
                    f"سریال موجود: {in_stock_serial_count}"
                )

        product.qty = new_qty

        if product.has_serial and not new_has_serial:
            serial_count = ProductSerial.query.filter_by(
                product_id=product.id
            ).count()

            if serial_count > 0:
                return (
                    "برای این کالا سریال ثبت شده است؛ "
                    "ابتدا باید سریال‌های ثبت‌شده را بررسی یا حذف کنید."
                )

        product.has_serial = new_has_serial

        db.session.commit()

        log_activity(
            session["user"],
            "EDIT_PRODUCT",
            f"ویرایش کالا: {old_name} ({old_qty}) -> {product.name} ({product.qty})"
        )

        return redirect("/dashboard")

    return render_template(
        "edit_product.html",
        product=product
    )

@bp.route("/export_excel")
def export_excel():

    if "user" not in session:
        return redirect("/")

    wb = Workbook()
    ws = wb.active

    ws.title = "Inventory"

    ws.append([
        "ID",
        "Product Name",
        "Quantity"
    ])

    products = Product.query.all()

    for p in products:
        ws.append([
            p.id,
            p.name,
            p.qty
        ])

    output = io.BytesIO()

    wb.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="inventory.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@bp.route("/users")
def users():

    if "user" not in session:
        return redirect("/")

    if session.get("role") != "admin":
        return "Access Denied"

    users = User.query.all()

    return render_template(
        "users.html",
        users=users
    )


@bp.route("/users/add", methods=["POST"])
def add_user():

    if session.get("role") != "admin":
        return "Access Denied"

    username = request.form["username"]
    role = request.form["role"]

    user = User(
        username=username,
        password=request.form["password"],
        role=role
    )

    db.session.add(user)
    db.session.commit()

    log_activity(
        session["user"],
        "ADD_USER",
        f"ایجاد کاربر: {username} | نقش: {role}"
    )

    return redirect("/users")

@bp.route("/reports")
def reports():

    if "user" not in session:
        return redirect("/")

    movement_type = request.args.get("type")
    created_by = request.args.get("created_by")
    product_name = request.args.get("product_name")

    date_from = request.args.get("date_from")
    date_to = request.args.get("date_to")

    query = Movement.query

    if movement_type:
        query = query.filter_by(
            type=movement_type
        )

    if created_by:
        query = query.filter(
            Movement.created_by.contains(created_by)
        )

    if date_from:
        query = query.filter(
            Movement.timestamp >= date_from
        )

    if date_to:
        query = query.filter(
            Movement.timestamp <= date_to + " 23:59:59"
        )

    if product_name:

        product_ids = [
            p.id
            for p in Product.query.filter(
                Product.name.contains(product_name)
            ).all()
        ]

        query = query.filter(
            Movement.product_id.in_(product_ids)
        )

    page = request.args.get("page", 1, type=int)

    pagination = query.order_by(
        Movement.id.desc()
    ).paginate(
        page=page,
        per_page=30,
        error_out=False
    )

    movements = pagination.items

    products_dict = {}

    for p in Product.query.all():
        products_dict[p.id] = p.name

    movement_serial_counts = {}

    for movement in movements:
        movement_serial_counts[movement.id] = ProductSerial.query.filter_by(
            movement_id=movement.id
        ).count()

    product_serial_flags = {}

    for product in Product.query.all():
        product_serial_flags[product.id] = product.has_serial

    return render_template(
        "reports.html",
        movements=movements,
        products_dict=products_dict,
        movement_serial_counts=movement_serial_counts,
        product_serial_flags=product_serial_flags,
        pagination=pagination
    )

@bp.route("/export_report_excel")
def export_report_excel():

    if "user" not in session:
        return redirect("/")

    wb = Workbook()
    ws = wb.active

    ws.title = "Reports"

    ws.append([
        "Product",
        "Type",
        "Quantity",
        "Receiver",
        "Customer",
        "Created By",
        "Timestamp"
    ])

    products_dict = {}

    for p in Product.query.all():
        products_dict[p.id] = p.name

    movements = Movement.query.order_by(
        Movement.id.desc()
    ).all()

    for m in movements:

        ws.append([
            products_dict.get(
                m.product_id,
                m.product_id
            ),
            m.type,
            m.qty,
            m.receiver_name,
            m.customer_name,
            m.created_by,
            str(m.timestamp)
        ])

    output = io.BytesIO()

    wb.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@bp.route("/backup")
def backup_database():

    if "user" not in session:
        return redirect("/")

    if session.get("role") != "admin":
        return "Access Denied"

    from datetime import datetime
    import shutil

    backup_dir = "/app/instance/backups"
    os.makedirs(backup_dir, exist_ok=True)

    filename = (
        "backup_inventory_"
        + datetime.now().strftime("%Y-%m-%d_%H%M%S")
        + ".db"
    )

    backup_path = os.path.join(
        backup_dir,
        filename
    )

    shutil.copy2(
        "/app/instance/data.db",
        backup_path
    )

    log_activity(
        session["user"],
        "BACKUP_DATABASE",
        f"تهیه نسخه پشتیبان: {filename}"
    )

    return send_file(
        backup_path,
        as_attachment=True,
        download_name=filename
    )

@bp.route("/activity-log")
def activity_log():

    if "user" not in session:
        return redirect("/")

    if session.get("role") != "admin":
        return "Access Denied"

    username = request.args.get("username", "").strip()
    action = request.args.get("action", "").strip()
    page = request.args.get("page", 1, type=int)

    query = ActivityLog.query

    if username:
        query = query.filter(
            ActivityLog.username.contains(username)
        )

    if action:
        query = query.filter_by(
            action=action
        )

    pagination = query.order_by(
        ActivityLog.id.desc()
    ).paginate(
        page=page,
        per_page=30,
        error_out=False
    )

    logs = pagination.items

    return render_template(
        "activity_log.html",
        logs=logs,
        pagination=pagination,
        username=username,
        selected_action=action
    )

@bp.route("/export_activity_log_excel")
def export_activity_log_excel():

    if "user" not in session:
        return redirect("/")

    if session.get("role") != "admin":
        return "Access Denied"

    username = request.args.get("username")
    action = request.args.get("action")

    query = ActivityLog.query

    if username:
        query = query.filter(
            ActivityLog.username.contains(username)
        )

    if action:
        query = query.filter_by(
            action=action
        )

    logs = query.order_by(
        ActivityLog.id.desc()
    ).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Activity Log"

    ws.append([
        "کاربر",
        "عملیات",
        "توضیحات",
        "زمان"
    ])

    action_names = {
        "ADD_PRODUCT": "ثبت کالا",
        "EDIT_PRODUCT": "ویرایش کالا",
        "DELETE_PRODUCT": "حذف کالا",
        "STOCK_IN": "ورود کالا به انبار",
        "STOCK_OUT": "خروج کالا از انبار",
        "ADD_USER": "ایجاد کاربر",
        "DELIVER_SERIAL": "تحویل سریال",
        "LOGIN": "ورود به سیستم",
        "LOGOUT": "خروج از سیستم"
    }

    for log in logs:
        ws.append([
            log.username,
            action_names.get(log.action, log.action),
            log.description,
            log.timestamp.strftime("%Y-%m-%d %H:%M:%S")
            if log.timestamp else ""
        ])

    from io import BytesIO

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="activity_log.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

@bp.route("/print-slip/<int:id>")
def print_slip(id):

    if "user" not in session:
        return redirect("/")

    movement = Movement.query.get(id)

    if not movement:
        return "حواله پیدا نشد"

    product = Product.query.get(
        movement.product_id
    )

    serials = ProductSerial.query.filter_by(
        movement_id=movement.id
    ).order_by(
        ProductSerial.id.asc()
    ).all()

    return render_template(
        "print_slip.html",
        movement=movement,
        product=product,
        serials=serials
    )

@bp.route("/print-inventory-report")
def print_inventory_report():

    if "user" not in session:
        return redirect("/")

    products = Product.query.order_by(
        Product.name
    ).all()

    report_data = []

    for product in products:

        last_in = Movement.query.filter_by(
            product_id=product.id,
            type="IN"
        ).order_by(
            Movement.id.desc()
        ).first()

        last_out = Movement.query.filter_by(
            product_id=product.id,
            type="OUT"
        ).order_by(
            Movement.id.desc()
        ).first()

        last_movement = Movement.query.filter_by(
            product_id=product.id
        ).order_by(
            Movement.id.desc()
        ).first()

        report_data.append({
            "product": product,
            "created_by": last_movement.created_by if last_movement else "",
            "last_in": last_in.timestamp if last_in else None,
            "last_out": last_out.timestamp if last_out else None
        })

    return render_template(
        "inventory_report.html",
        report_data=report_data
    )

@bp.route("/serials/<int:product_id>")
def serials(product_id):

    if "user" not in session:
        return redirect("/")

    product = Product.query.get(product_id)

    if not product:
        return "کالا پیدا نشد"

    if not product.has_serial:
        return "این کالا شماره سریال ندارد"

    serials = ProductSerial.query.filter_by(
        product_id=product_id
    ).order_by(
        ProductSerial.id.desc()
    ).all()

    out_movements = Movement.query.filter_by(
        product_id=product_id,
        type="OUT"
    ).order_by(
        Movement.id.desc()
    ).all()

    movement_serial_counts = {}

    for movement in out_movements:
        movement_serial_counts[movement.id] = ProductSerial.query.filter_by(
            movement_id=movement.id
        ).count()

    return render_template(
        "serials.html",
        product=product,
        serials=serials,
        out_movements=out_movements,
        movement_serial_counts=movement_serial_counts
    )

@bp.route("/serials/add", methods=["POST"])
def add_serial():

    if "user" not in session:
        return redirect("/")

    product_id = request.form["product_id"]
    serial_number = request.form["serial_number"].strip()

    if not serial_number:
        return "شماره سریال وارد نشده است"

    exists = ProductSerial.query.filter_by(
        serial_number=serial_number
    ).first()

    if exists:
        return "این سریال قبلاً ثبت شده است"

    serial = ProductSerial(
        product_id=product_id,
        serial_number=serial_number,
        created_by=session["user"]
    )

    db.session.add(serial)
    db.session.commit()

    return redirect(
        f"/serials/{product_id}"
    )

@bp.route("/serial-search")
def serial_search():

    if "user" not in session:
        return redirect("/")

    serial_number = request.args.get(
        "serial_number",
        ""
    )

    result = None
    product = None

    if serial_number:

        result = ProductSerial.query.filter(
            ProductSerial.serial_number.contains(
                serial_number
            )
        ).first()

        if result:
            product = Product.query.get(
                result.product_id
            )

    return render_template(
        "serial_search.html",
        result=result,
        product=product
    )

@bp.route("/serials/deliver/<int:id>", methods=["POST"])
def deliver_serial(id):

    if "user" not in session:
        return redirect("/")

    serial = ProductSerial.query.get(id)

    if not serial:
        return "سریال پیدا نشد"

    movement_id = request.form.get("movement_id", type=int)

    if not movement_id:
        return "شماره حواله انتخاب نشده است"

    movement = Movement.query.get(movement_id)

    if not movement:
        return "حواله پیدا نشد"

    if movement.type != "OUT":
        return "فقط حواله خروج قابل انتخاب است"

    if movement.product_id != serial.product_id:
        return "این حواله مربوط به کالای دیگری است"

    delivered_count = ProductSerial.query.filter_by(
        movement_id=movement.id
    ).count()

    if delivered_count >= movement.qty:
        return "تعداد سریال‌های تحویل‌شده برای این حواله کامل شده است"

    serial.status = "DELIVERED"
    serial.movement_id = movement.id
    serial.customer_name = movement.customer_name

    db.session.commit()

    log_activity(
        session["user"],
        "DELIVER_SERIAL",
        f"تحویل سریال {serial.serial_number} | کالا: {serial.product_id} | حواله خروج: {movement.id} | مشتری: {movement.customer_name or '-'}"
    )

    return redirect(
        f"/serials/{serial.product_id}"
    )

@bp.route("/serials/delete/<int:id>")
def delete_serial(id):

    if "user" not in session:
        return redirect("/")

    if session.get("role") != "admin":
        return "شما اجازه حذف سریال را ندارید"

    serial = ProductSerial.query.get(id)

    if not serial:
        return "سریال پیدا نشد"

    if serial.status != "IN_STOCK":
        if serial.movement_id:
            return (
                f"این سریال به حواله خروج شماره {serial.movement_id} "
                "وصل شده است و برای حفظ سابقه قابل حذف نیست"
            )

        return "این سریال تحویل شده است و برای حفظ سابقه قابل حذف نیست"

    product_id = serial.product_id

    db.session.delete(serial)
    db.session.commit()

    return redirect(f"/serials/{product_id}")
