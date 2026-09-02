
from flask import Blueprint, render_template, request, redirect, session, flash
from app.models import (
    db,
    Product,
    Movement,
    User,
    ActivityLog,
    ProductSerial,
    MaterialRequest,
    MaterialRequestItem,
    Project,
    RequestNote,
    InventoryDocument,
    InventoryDocumentItem,
    ProjectReturn
)
from openpyxl import Workbook
from flask import send_file, jsonify
from sqlalchemy import func
import io
import os
import jdatetime
from datetime import datetime

def normalize_digits(value):

    if value is None:
        return ""

    persian = "۰۱۲۳۴۵۶۷۸۹"
    arabic = "٠١٢٣٤٥٦٧٨٩"
    english = "0123456789"

    for i in range(10):
        value = value.replace(persian[i], english[i])
        value = value.replace(arabic[i], english[i])

    return value

bp = Blueprint("main", __name__)

@bp.app_template_filter("jalali")
def jalali_date(value):

    if not value:
        return ""

    return jdatetime.datetime.fromgregorian(
        datetime=value
    ).strftime("%Y/%m/%d %H:%M")

def log_activity(username, action, description):

    try:

        log = ActivityLog(
            username=username,
            action=action,
            description=description
        )

        db.session.add(log)
        db.session.commit()

    except Exception:
        db.session.rollback()

def shamsi_date(value):

    if not value:
        return ""

    return jdatetime.datetime.fromgregorian(
        datetime=value
    ).strftime("%Y/%m/%d %H:%M")

bp.add_app_template_filter(
    shamsi_date,
    "shamsi"
)

@bp.route("/")
def login():
    return render_template("login.html")

@bp.route("/login", methods=["POST"])
def do_login():

    username = request.form["username"]
    password = request.form["password"]

    user = User.query.filter_by(
        username=username,
        password=password
    ).first()

    if user:

        session["user"] = user.username
        session["role"] = user.role

        log_activity(
            user.username,
            "LOGIN",
            "ورود به سیستم"
        )

        return redirect("/dashboard")

    return "نام کاربری یا رمز عبور اشتباه است"

@bp.route("/dashboard")
def dashboard():

    if "user" not in session:
        return redirect("/")

    search = request.args.get("search")

    all_products_query = Product.query.order_by(
        Product.id.desc()
    )

    if search:
        all_products_query = all_products_query.filter(
            Product.name.contains(search)
        )

    all_products = all_products_query.all()

    # فقط 10 کالا در داشبورد نمایش داده می‌شود
    products = all_products[:10]

    all_active_products = Product.query.filter_by(
        is_active=1
    ).order_by(
        Product.name
    ).all()

    movements = Movement.query.order_by(
        Movement.id.desc()
    ).limit(10).all()

    total_products = len(all_products)

    low_stock = len([
        p for p in all_products
        if p.qty > 0 and p.qty <= p.min_qty
    ])

    out_of_stock = len([
        p for p in all_products
        if p.qty == 0
    ])

    pending_requests = MaterialRequest.query.filter_by(
        status="PENDING"
    ).count()

    approved_requests = MaterialRequest.query.filter_by(
        status="APPROVED"
    ).count()

    low_stock_products = [
        p for p in Product.query.filter(
            Product.is_active == 1,
            Product.qty > 0
        ).all()
        if p.qty <= p.min_qty
    ]

    low_stock_products = sorted(
        low_stock_products,
        key=lambda p: (p.qty, p.name)
    )[:8]

    out_of_stock_products = Product.query.filter(
        Product.is_active == 1,
        Product.qty == 0
    ).order_by(
        Product.name.asc()
    ).limit(8).all()

    serial_counts = {}

    serial_counts = {}

    for p in all_products:

        serial_counts[p.id] = ProductSerial.query.filter_by(
            product_id=p.id,
            status="IN_STOCK"
        ).count()

    products_dict = {}
    products_units = {}

    for product in Product.query.all():
        products_dict[product.id] = product.name
        products_units[product.id] = product.unit or "عدد"

    return render_template(
        "dashboard.html",
        products=products,
        movements=movements,
        total_products=total_products,
        low_stock=low_stock,
        out_of_stock=out_of_stock,
        serial_counts=serial_counts,
        products_dict=products_dict,
        products_units=products_units,
        all_active_products=all_active_products,
        low_stock_products=low_stock_products,
        out_of_stock_products=out_of_stock_products,
        pending_requests=pending_requests,
        approved_requests=approved_requests,
    )

@bp.route("/products")
def products_page():

    if "user" not in session:
        return redirect("/")

    search = request.args.get("search", "").strip()
    category = request.args.get("category", "").strip()
    stock_status = request.args.get("stock_status", "").strip()
    condition = request.args.get("condition", "").strip()
    active = request.args.get("active", "").strip()
    page = request.args.get("page", 1, type=int)

    query = Product.query.order_by(
        Product.is_active.desc(),
        Product.id.desc()
    )

    if search:
        query = query.filter(Product.name.contains(search))

    if category:
        query = query.filter(Product.category == category)

    if condition:
        query = query.filter(Product.condition == condition)

    if active == "1":
       query = query.filter(Product.is_active == 1)

    elif active == "0":
        query = query.filter(Product.is_active == 0)

    if stock_status == "in_stock":
        query = query.filter(Product.qty > Product.min_qty)

    elif stock_status == "low_stock":
        query = query.filter(
            Product.qty > 0,
            Product.qty <= Product.min_qty
        )

    elif stock_status == "out_of_stock":
        query = query.filter(Product.qty == 0)

    pagination = query.paginate(
        page=page,
        per_page=20,
        error_out=False
    )

    serial_counts = {}

    for product in pagination.items:
        serial_counts[product.id] = ProductSerial.query.filter_by(
            product_id=product.id,
            status="IN_STOCK"
        ).count()

    categories = [
        "دوربین مداربسته",
        "دستگاه ضبط NVR / DVR",
        "هارد دیسک",
        "تجهیزات شبکه",
        "تجهیزات مخابراتی",
        "نمایشگر",
        "قطعات کامپیوتر",
        "کابل و متعلقات",
        "لوازم نصب",
        "منبع تغذیه و برق",
        "سایر"
    ]

    return render_template(
        "products.html",
        products=pagination.items,
        pagination=pagination,
        search=search,
        category=category,
        stock_status=stock_status,
        condition=condition,
        active=active,
        categories=categories,
        serial_counts=serial_counts
    )

@bp.route("/add", methods=["POST"])
def add():

    if "user" not in session:
        return redirect("/")

    if session.get("role") == "viewer":
        flash("شما اجازه ثبت کالا ندارید.", "danger")
        return redirect("/products")

    name = request.form.get("name", "").strip()
    qty_text = normalize_digits(
        request.form.get("qty", "0").strip()
    )

    min_qty_text = normalize_digits(
        request.form.get("min_qty", "5").strip()
    )
    unit = request.form.get("unit", "عدد").strip()
    category = request.form.get("category", "سایر").strip() or "سایر"
    condition = request.form.get("condition", "نو").strip() or "نو"
    has_serial = int(request.form.get("has_serial", 0))

    if not name:
        flash("نام کالا وارد نشده است.", "danger")
        return redirect("/products")

    try:
        qty = int(qty_text)
        min_qty = int(min_qty_text)
    except ValueError:
        flash("موجودی اولیه و حداقل موجودی باید عدد باشند.", "danger")
        return redirect("/products")

    if qty < 0 or min_qty < 0:
        flash("موجودی اولیه و حداقل موجودی نمی‌توانند منفی باشند.", "danger")
        return redirect("/products")

    normalized_name = (
        name
        .replace("ي", "ی")
        .replace("ك", "ک")
        .replace("\u200c", " ")
        .strip()
        .lower()
    )

    for item in Product.query.all():
        existing_name = (
            (item.name or "")
            .replace("ي", "ی")
            .replace("ك", "ک")
            .replace("\u200c", " ")
            .strip()
            .lower()
        )

        if existing_name == normalized_name:
            flash(f"کالایی با نام «{name}» قبلاً ثبت شده است.", "warning")
            return redirect("/products")

    p = Product(
        name=name,
        qty=qty,
        has_serial=has_serial,
        min_qty=min_qty,
        unit=unit or "عدد",
        category=category,
        condition=condition
    )

    db.session.add(p)
    db.session.commit()

    log_activity(
        session["user"],
        "ADD_PRODUCT",
        f"افزودن کالا: {name} | موجودی اولیه: {qty} | حداقل موجودی: {min_qty}"
    )

    flash(f"کالای «{name}» با موفقیت ثبت شد.", "success")
    return redirect("/products")

def build_movement_query():

    movement_type = request.args.get("type", "").strip()
    created_by = request.args.get("created_by", "").strip()
    product_name = request.args.get("product_name", "").strip()
    project_name = request.args.get("project_name", "").strip()

    date_from = request.args.get("date_from", "").strip()
    date_to = request.args.get("date_to", "").strip()

    query = Movement.query

    if movement_type:
        query = query.filter(
            Movement.type == movement_type
        )

    if created_by:
        query = query.filter(
            Movement.created_by.contains(created_by)
        )

    if project_name:
        query = query.filter(
            Movement.project_name.contains(project_name)
        )

    if product_name:

        product_ids = [
            product.id
            for product in Product.query.filter(
                Product.name.contains(product_name)
            ).all()
        ]

        query = query.filter(
            Movement.product_id.in_(product_ids)
        )

    if date_from:

        y, m, d = map(int, date_from.split("/"))

        g_from = (
            jdatetime.date(y, m, d)
            .togregorian()
        )

        query = query.filter(
            Movement.timestamp >= datetime.combine(
                g_from,
                datetime.min.time()
            )
        )

    if date_to:

        y, m, d = map(int, date_to.split("/"))

        g_to = (
            jdatetime.date(y, m, d)
            .togregorian()
        )

        query = query.filter(
            Movement.timestamp <= datetime.combine(
                g_to,
                datetime.max.time()
            )
        )

    return query

@bp.route("/move", methods=["POST"])
def move():

    product_id = int(request.form["product_id"])
    qty = int(request.form["qty"])
    mtype = request.form["type"]
    receiver_name = request.form.get(
        "receiver_name",
        ""
    )

    customer_name = request.form.get(
        "customer_name",
        ""
    )

    project_name = request.form.get(
        "project_name",
        ""
    ).strip()

    description = request.form.get(
        "description",
        ""
    )

    product = Product.query.get(product_id)

    if mtype == "OUT" and product.qty < qty:
        flash("موجودی کالا برای این خروج کافی نیست.", "danger")
        return redirect("/products")

    if mtype == "OUT" and product.has_serial:
        available_serials = ProductSerial.query.filter_by(
            product_id=product.id,
            status="IN_STOCK"
        ).count()

        if available_serials < qty:
            flash(
                f"تعداد سریال‌های موجود کافی نیست. موجود: {available_serials} - درخواست خروج: {qty}",
                "danger"
            )
            return redirect("/products")

    if mtype == "IN":
        product.qty += qty
    else:
        product.qty -= qty

    m = Movement(
        product_id=product_id,
        qty=qty,
        type=mtype,
        receiver_name=receiver_name,
        customer_name=customer_name,
        project_name=project_name,
        description=description,
        created_by=session["user"]
    )
    db.session.add(m)
    db.session.commit()

    if mtype == "IN":

        log_activity(
            session["user"],
            "STOCK_IN",
            f"ورود {qty} عدد از کالا {product.name}"
        )

    else:

        log_activity(
            session["user"],
            "STOCK_OUT",
            f"خروج {qty} عدد از کالا {product.name}"
        )

    action_text = "ورود به انبار" if mtype == "IN" else "خروج از انبار"
    flash(
        f"{action_text} برای کالای «{product.name}» با تعداد {qty} با موفقیت ثبت شد.",
        "success"
    )

    return redirect("/products")
@bp.route("/logout")
def logout():

    if "user" in session:

        log_activity(
            session["user"],
            "LOGOUT",
            "خروج از سیستم"
        )

    session.clear()

    return redirect("/")

@bp.route("/delete/<int:id>")
def delete_product(id):

    if "user" not in session:
        return redirect("/")

    if session.get("role") != "admin":
        flash("فقط مدیر سیستم اجازه حذف کالا را دارد.", "danger")
        return redirect("/products")

    product = Product.query.get(id)

    if not product:
        flash("کالا پیدا نشد.", "warning")
        return redirect("/products")

    movement_count = Movement.query.filter_by(
        product_id=product.id
    ).count()

    serial_count = ProductSerial.query.filter_by(
        product_id=product.id
    ).count()

    if movement_count > 0 or serial_count > 0:
        flash(
            f"کالای «{product.name}» حذف نشد؛ زیرا دارای "
            f"{movement_count} گردش و {serial_count} شماره سریال ثبت‌شده است.",
            "warning"
        )
        return redirect("/products")

    product_name = product.name

    db.session.delete(product)
    db.session.commit()

    log_activity(
        session["user"],
        "DELETE_PRODUCT",
        f"حذف کالا: {product_name}"
    )

    flash(f"کالای «{product_name}» با موفقیت حذف شد.", "success")
    return redirect("/products")

@bp.route("/product-toggle-active/<int:id>")
def product_toggle_active(id):

    if "user" not in session:
        return redirect("/")

    if session.get("role") != "admin":
        flash("فقط مدیر سیستم اجازه تغییر وضعیت کالا را دارد.", "danger")
        return redirect("/products")

    product = Product.query.get(id)

    if not product:
        flash("کالا پیدا نشد.", "warning")
        return redirect("/products")

    product.is_active = 0 if product.is_active else 1
    db.session.commit()

    status_text = "فعال شد" if product.is_active else "غیرفعال شد"

    log_activity(
        session["user"],
        "TOGGLE_PRODUCT_ACTIVE",
        f"وضعیت کالا «{product.name}»: {status_text}"
    )

    flash(f"کالای «{product.name}» {status_text}.", "success")
    return redirect("/products")

@bp.route("/edit/<int:id>", methods=["GET","POST"])
def edit_product(id):

    if "user" not in session:
        return redirect("/")

    if session.get("role") == "viewer":
        flash("شما اجازه ویرایش کالا ندارید.", "danger")
        return redirect("/products")

    product = Product.query.get(id)

    if not product:
        flash("کالا پیدا نشد.", "warning")
        return redirect("/products")

    if request.method == "POST":

        new_name = request.form.get("name", "").strip()
        qty_text = request.form.get("qty", "").strip()
        min_qty_text = request.form.get("min_qty", "5").strip()
        new_unit = request.form.get("unit", "عدد").strip() or "عدد"
        new_category = request.form.get("category", "سایر").strip() or "سایر"
        new_condition = request.form.get("condition", "نو").strip() or "نو"
        new_has_serial = 1 if request.form.get("has_serial") == "1" else 0

        if not new_name:
            flash("نام کالا وارد نشده است.", "danger")
            return redirect(f"/edit/{product.id}")

        try:
            new_qty = int(qty_text)
            new_min_qty = int(min_qty_text)
        except ValueError:
            flash("موجودی و حداقل موجودی باید عدد باشند.", "danger")
            return redirect(f"/edit/{product.id}")

        if new_qty < 0 or new_min_qty < 0:
            flash("موجودی و حداقل موجودی نمی‌توانند منفی باشند.", "danger")
            return redirect(f"/edit/{product.id}")

        normalized_name = (
            new_name
            .replace("ي", "ی")
            .replace("ك", "ک")
            .replace("\u200c", " ")
            .strip()
            .lower()
        )

        for item in Product.query.filter(Product.id != product.id).all():
            existing_name = (
                (item.name or "")
                .replace("ي", "ی")
                .replace("ك", "ک")
                .replace("\u200c", " ")
                .strip()
                .lower()
            )

            if existing_name == normalized_name:
                flash(f"کالایی با نام «{new_name}» قبلاً ثبت شده است.", "warning")
                return redirect(f"/edit/{product.id}")

        if new_has_serial:
            in_stock_serial_count = ProductSerial.query.filter_by(
                product_id=product.id,
                status="IN_STOCK"
            ).count()

            if new_qty < in_stock_serial_count:
                flash(
                    f"موجودی جدید نمی‌تواند کمتر از تعداد سریال‌های موجود باشد. "
                    f"سریال موجود: {in_stock_serial_count}",
                    "danger"
                )
                return redirect(f"/edit/{product.id}")

        if product.has_serial and not new_has_serial:
            serial_count = ProductSerial.query.filter_by(
                product_id=product.id
            ).count()

            if serial_count > 0:
                flash(
                    "برای این کالا سریال ثبت شده است؛ ابتدا باید سریال‌های ثبت‌شده را بررسی یا حذف کنید.",
                    "danger"
                )
                return redirect(f"/edit/{product.id}")

        old_name = product.name
        old_qty = product.qty
        old_min_qty = product.min_qty

        product.name = new_name
        product.qty = new_qty
        product.min_qty = new_min_qty
        product.unit = new_unit
        product.category = new_category
        product.condition = new_condition
        product.has_serial = new_has_serial

        db.session.commit()

        log_activity(
            session["user"],
            "EDIT_PRODUCT",
            f"ویرایش کالا: {old_name} ({old_qty}) -> {product.name} ({product.qty}) | حداقل موجودی: {old_min_qty} -> {product.min_qty}"
        )

        flash(f"تغییرات کالای «{product.name}» با موفقیت ذخیره شد.", "success")
        return redirect("/products")

    return render_template(
        "edit_product.html",
        product=product
    )

@bp.route("/export_excel")
def export_excel():

    if "user" not in session:
        return redirect("/")

    wb = Workbook()
    ws = wb.active

    ws.title = "Inventory"

    ws.append([
        "ID",
        "Product Name",
        "Quantity"
    ])

    products = Product.query.all()

    for p in products:
        ws.append([
            p.id,
            p.name,
            p.qty
        ])

    output = io.BytesIO()

    wb.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="inventory.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@bp.route("/users")
def users():

    if "user" not in session:
        return redirect("/")

    if session.get("role") != "admin":
        return "Access Denied"

    users = User.query.all()

    return render_template(
        "users.html",
        users=users
    )


@bp.route("/users/add", methods=["POST"])
def add_user():

    if session.get("role") != "admin":
        return "Access Denied"

    username = request.form["username"]
    role = request.form["role"]

    user = User(
        username=username,
        password=request.form["password"],
        role=role
    )

    db.session.add(user)
    db.session.commit()

    log_activity(
        session["user"],
        "ADD_USER",
        f"ایجاد کاربر: {username} | نقش: {role}"
    )

    return redirect("/users")

@bp.route("/users/delete/<int:id>")
def delete_user(id):

    if "user" not in session:
        return redirect("/")

    if session.get("role") != "admin":
        return "Access Denied"

    user = User.query.get(id)

    if not user:
        return "کاربر پیدا نشد"

    # جلوگیری از حذف کاربر فعلی
    if user.username == session["user"]:
        flash(
            "شما نمی‌توانید کاربری که با آن وارد شده‌اید را حذف کنید.",
            "danger"
        )
        return redirect("/users")

    # جلوگیری از حذف آخرین مدیر
    if user.role == "admin":

        admins = User.query.filter_by(role="admin").count()

        if admins <= 1:

            flash(
                "آخرین مدیر سیستم قابل حذف نیست.",
                "danger"
            )

            return redirect("/users")

    username = user.username

    db.session.delete(user)
    db.session.commit()

    log_activity(
        session["user"],
        "DELETE_USER",
        f"حذف کاربر: {username}"
    )

    flash(
        f"کاربر {username} حذف شد.",
        "success"
    )

    return redirect("/users")

@bp.route("/users/edit/<int:id>")
def edit_user(id):

    if "user" not in session:
        return redirect("/")

    if session.get("role") != "admin":
        return "Access Denied"

    user = User.query.get(id)

    if not user:
        return "کاربر پیدا نشد"

    return render_template(
        "edit_user.html",
        user=user
    )

@bp.route("/users/update/<int:id>", methods=["POST"])
def update_user(id):

    if "user" not in session:
        return redirect("/")

    if session.get("role") != "admin":
        return "Access Denied"

    user = User.query.get(id)

    if not user:
        return "کاربر پیدا نشد"

    username = request.form["username"].strip()
    role = request.form["role"]
    password = request.form["password"].strip()

    user.username = username
    user.role = role

    if password:
        user.password = password

    db.session.commit()

    log_activity(
        session["user"],
        "EDIT_USER",
        f"ویرایش کاربر: {username}"
    )

    flash(
        "اطلاعات کاربر با موفقیت ذخیره شد.",
        "success"
    )

    return redirect("/users")

@bp.route("/reports")
def reports():

    if "user" not in session:
        return redirect("/")

    query = build_movement_query()

    page = request.args.get("page", 1, type=int)

    per_page = request.args.get("per_page", 10, type=int)

    if per_page not in (10, 25, 50, 100):
        per_page = 10

    pagination = query.order_by(
        Movement.id.desc()
    ).paginate(
        page=page,
        per_page=per_page,
        error_out=False
    )

    movements = pagination.items

    products_dict = {}
    products_units = {}

    for p in Product.query.all():
        products_dict[p.id] = p.name
        products_units[p.id] = p.unit or "عدد"

    product_serial_flags = {}

    for p in Product.query.all():
        product_serial_flags[p.id] = 1 if p.has_serial else 0

    movement_serial_counts = {}

    for movement in movements:
        movement_serial_counts[movement.id] = ProductSerial.query.filter_by(
            movement_id=movement.id
        ).count()

    return render_template(
        "reports.html",
        movements=movements,
        products_dict=products_dict,
        products_units=products_units,
        product_serial_flags=product_serial_flags,
        movement_serial_counts=movement_serial_counts,
        pagination=pagination
    )

@bp.route("/export_report_excel")
def export_report_excel():

    if "user" not in session:
        return redirect("/")

    query = build_movement_query()

    movements = (
        query
        .order_by(Movement.id.desc())
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Reports"

    ws.append([
        "نام کالا",
        "نوع عملیات",
        "تعداد",
        "تحویل گیرنده",
        "مشتری",
        "پروژه",
        "شماره سریال",
        "ثبت کننده",
        "تاریخ"
    ])

    products_dict = {}

    for product in Product.query.all():
        products_dict[product.id] = product.name

    for movement in movements:

        serials = (
            ProductSerial.query
            .filter_by(movement_id=movement.id)
            .order_by(ProductSerial.id.asc())
            .all()
        )

        serial_numbers = ", ".join(
            serial.serial_number
            for serial in serials
            if serial.serial_number
        )

        ws.append([

            products_dict.get(
                movement.product_id,
                movement.product_id
            ),

            "ورود" if movement.type == "IN" else "خروج",

            movement.qty,

            movement.receiver_name or "",

            movement.customer_name or "",

            movement.project_name or "",

            serial_numbers,

            movement.created_by or "",

            jdatetime.datetime.fromgregorian(
                datetime=movement.timestamp
            ).strftime("%Y/%m/%d %H:%M")

        ])

    output = io.BytesIO()

    wb.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="report.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

@bp.route("/backup")
def backup_database():

    if "user" not in session:
        return redirect("/")

    if session.get("role") != "admin":
        return "Access Denied"

    from datetime import datetime
    import shutil

    backup_dir = "/app/instance/backups"
    os.makedirs(backup_dir, exist_ok=True)

    filename = (
        "backup_inventory_"
        + datetime.now().strftime("%Y-%m-%d_%H%M%S")
        + ".db"
    )

    backup_path = os.path.join(
        backup_dir,
        filename
    )

    shutil.copy2(
        "/app/instance/data.db",
        backup_path
    )

    log_activity(
        session["user"],
        "BACKUP_DATABASE",
        f"تهیه نسخه پشتیبان: {filename}"
    )

    return send_file(
        backup_path,
        as_attachment=True,
        download_name=filename
    )

@bp.route("/activity-log")
def activity_log():

    if "user" not in session:
        return redirect("/")

    if session.get("role") != "admin":
        return "Access Denied"

    username = request.args.get("username", "").strip()
    action = request.args.get("action", "").strip()
    page = request.args.get("page", 1, type=int)

    query = ActivityLog.query

    if username:
        query = query.filter(
            ActivityLog.username.contains(username)
        )

    if action:
        query = query.filter_by(
            action=action
        )

    pagination = query.order_by(
        ActivityLog.id.desc()
    ).paginate(
        page=page,
        per_page=10,
        error_out=False
    )

    logs = pagination.items

    return render_template(
        "activity_log.html",
        logs=logs,
        pagination=pagination,
        username=username,
        selected_action=action
    )

@bp.route("/export_activity_log_excel")
def export_activity_log_excel():

    if "user" not in session:
        return redirect("/")

    if session.get("role") != "admin":
        return "Access Denied"

    username = request.args.get("username")
    action = request.args.get("action")

    query = ActivityLog.query

    if username:
        query = query.filter(
            ActivityLog.username.contains(username)
        )

    if action:
        query = query.filter_by(
            action=action
        )

    logs = query.order_by(
        ActivityLog.id.desc()
    ).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Activity Log"

    ws.append([
        "کاربر",
        "عملیات",
        "توضیحات",
        "زمان"
    ])

    action_names = {
        "ADD_PRODUCT": "ثبت کالا",
        "EDIT_PRODUCT": "ویرایش کالا",
        "DELETE_PRODUCT": "حذف کالا",
        "STOCK_IN": "ورود کالا به انبار",
        "STOCK_OUT": "خروج کالا از انبار",
        "ADD_USER": "ایجاد کاربر",
        "DELIVER_SERIAL": "تحویل سریال",
        "LOGIN": "ورود به سیستم",
        "LOGOUT": "خروج از سیستم"
    }

    for log in logs:
        ws.append([
            log.username,
            action_names.get(log.action, log.action),
            log.description,
            log.timestamp.strftime("%Y-%m-%d %H:%M:%S")
            if log.timestamp else ""
        ])

    from io import BytesIO

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="activity_log.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

@bp.route("/print-slip/<int:id>")
def print_slip(id):

    if "user" not in session:
        return redirect("/")

    movement = Movement.query.get(id)

    if not movement:
        return "حواله پیدا نشد"

    product = Product.query.get(
        movement.product_id
    )

    serials = ProductSerial.query.filter_by(
        movement_id=movement.id
    ).order_by(
        ProductSerial.id.asc()
    ).all()

    return render_template(
        "print_slip.html",
        movement=movement,
        product=product,
        serials=serials
    )


@bp.route("/inventory")
def inventory():

    if "user" not in session:
        return redirect("/")

    movements = (
        Movement.query
        .order_by(Movement.timestamp.desc())
        .all()
    )

    products = {
        p.id: p
        for p in Product.query.all()
    }

    return render_template(
        "inventory.html",
        movements=movements,
        products=products
    )

@bp.route("/print-inventory-report")
def print_inventory_report():

    if "user" not in session:
        return redirect("/")

    search = request.args.get("search", "").strip()
    category = request.args.get("category", "").strip()
    stock_status = request.args.get("stock_status", "").strip()
    condition = request.args.get("condition", "").strip()
    active = request.args.get("active", "").strip()

    query = Product.query

    if search:
        query = query.filter(Product.name.contains(search))

    if category:
        query = query.filter(Product.category == category)

    if condition:
        query = query.filter(Product.condition == condition)

    if active == "1":
        query = query.filter(Product.is_active == 1)

    elif active == "0":
        query = query.filter(Product.is_active == 0)

    if stock_status == "in_stock":
        query = query.filter(Product.qty > Product.min_qty)

    elif stock_status == "low_stock":
        query = query.filter(
            Product.qty > 0,
            Product.qty <= Product.min_qty
        )

    elif stock_status == "out_of_stock":
        query = query.filter(Product.qty == 0)

    products = query.order_by(Product.name).all()

    return render_template(
        "print_inventory_report.html",
        products=products,
        search=search,
        category=category,
        stock_status=stock_status,
        condition=condition
    )


@bp.route("/print-movements-report")
def print_movements_report():

    if "user" not in session:
        return redirect("/")

    movement_type = request.args.get("type", "").strip()
    created_by = request.args.get("created_by", "").strip()
    product_name = request.args.get("product_name", "").strip()
    project_name = request.args.get("project_name", "").strip()
    date_from = request.args.get("date_from", "").strip()
    date_to = request.args.get("date_to", "").strip()

    query = Movement.query

    if movement_type:
        query = query.filter(Movement.type == movement_type)

    if created_by:
        query = query.filter(Movement.created_by.contains(created_by))

    if project_name:
        query = query.filter(Movement.project_name.contains(project_name))

    if date_from:
        query = query.filter(Movement.timestamp >= date_from)

    if date_to:
        query = query.filter(Movement.timestamp <= date_to + " 23:59:59")

    if product_name:
        product_ids = [
            product.id
            for product in Product.query.filter(
                Product.name.contains(product_name)
            ).all()
        ]
        query = query.filter(Movement.product_id.in_(product_ids))

    movements = query.order_by(Movement.id.desc()).all()

    products_dict = {}
    products_units = {}

    for product in Product.query.all():
        products_dict[product.id] = product.name
        products_units[product.id] = product.unit or "عدد"

    return render_template(
        "print_movements_report.html",
        movements=movements,
        products_dict=products_dict,
        products_units=products_units,
        movement_type=movement_type,
        created_by=created_by,
        product_name=product_name,
        project_name=project_name,
        date_from=date_from,
        date_to=date_to
    )

@bp.route("/serials")
def serials_index():

    if "user" not in session:
        return redirect("/")

    products = (
        Product.query
        .filter_by(
            has_serial=1,
            is_active=1
        )
        .order_by(Product.name)
        .all()
    )

    serial_counts = {}

    for product in products:

        serial_counts[product.id] = (
            ProductSerial.query
            .filter_by(product_id=product.id)
            .count()
        )

    return render_template(
        "serials_index.html",
        products=products,
        serial_counts=serial_counts
    )

@bp.route("/serials/<int:product_id>")
def serials(product_id):

    if "user" not in session:
        return redirect("/")

    product = Product.query.get(product_id)

    if not product:
        return "کالا پیدا نشد"

    if not product.has_serial:
        return "این کالا شماره سریال ندارد"

    serials = ProductSerial.query.filter_by(
        product_id=product_id
    ).order_by(
        ProductSerial.id.desc()
    ).all()

    out_movements = Movement.query.filter_by(
        product_id=product_id,
        type="OUT"
    ).order_by(
        Movement.id.desc()
    ).all()

    movement_serial_counts = {}

    for movement in out_movements:
        movement_serial_counts[movement.id] = ProductSerial.query.filter_by(
            movement_id=movement.id
        ).count()

    return render_template(
        "serials.html",
        product=product,
        serials=serials,
        out_movements=out_movements,
        movement_serial_counts=movement_serial_counts
    )

@bp.route("/serials/add", methods=["POST"])
def add_serial():

    if "user" not in session:
        return redirect("/")

    product_id = request.form["product_id"]
    serial_number = request.form["serial_number"].strip()

    if not serial_number:
        return "شماره سریال وارد نشده است"

    exists = ProductSerial.query.filter_by(
        serial_number=serial_number
    ).first()

    if exists:
        return "این سریال قبلاً ثبت شده است"

    serial = ProductSerial(
        product_id=product_id,
        serial_number=serial_number,
        created_by=session["user"]
    )

    db.session.add(serial)
    db.session.commit()

    return redirect(
        f"/serials/{product_id}"
    )

@bp.route("/serial-search")
def serial_search():

    if "user" not in session:
        return redirect("/")

    serial_number = request.args.get(
        "serial_number",
        ""
    )

    result = None
    product = None

    if serial_number:

        result = ProductSerial.query.filter(
            ProductSerial.serial_number.contains(
                serial_number
            )
        ).first()

        if result:
            product = Product.query.get(
                result.product_id
            )

    movement = None

    if result and result.movement_id:
        movement = Movement.query.get(result.movement_id)

    return render_template(
        "serial_search.html",
        result=result,
        product=product,
        movement=movement
    )

@bp.route("/serials/deliver/<int:id>", methods=["POST"])
def deliver_serial(id):

    if "user" not in session:
        return redirect("/")

    serial = ProductSerial.query.get(id)

    if not serial:
        return "سریال پیدا نشد"

    movement_id = request.form.get("movement_id", type=int)

    if not movement_id:
        return "شماره حواله انتخاب نشده است"

    movement = Movement.query.get(movement_id)

    if not movement:
        return "حواله پیدا نشد"

    if movement.type != "OUT":
        return "فقط حواله خروج قابل انتخاب است"

    if movement.product_id != serial.product_id:
        return "این حواله مربوط به کالای دیگری است"

    delivered_count = ProductSerial.query.filter_by(
        movement_id=movement.id
    ).count()

    if delivered_count >= movement.qty:
        return "تعداد سریال‌های تحویل‌شده برای این حواله کامل شده است"

    serial.status = "DELIVERED"
    serial.movement_id = movement.id
    serial.customer_name = movement.customer_name

    db.session.commit()

    log_activity(
        session["user"],
        "DELIVER_SERIAL",
        f"تحویل سریال {serial.serial_number} | کالا: {serial.product_id} | حواله خروج: {movement.id} | مشتری: {movement.customer_name or '-'}"
    )

    return redirect(
        f"/serials/{serial.product_id}"
    )

@bp.route("/serials/delete/<int:id>")
def delete_serial(id):

    if "user" not in session:
        return redirect("/")

    if session.get("role") != "admin":
        return "شما اجازه حذف سریال را ندارید"

    serial = ProductSerial.query.get(id)

    if not serial:
        return "سریال پیدا نشد"

    if serial.status != "IN_STOCK":
        if serial.movement_id:
            return (
                f"این سریال به حواله خروج شماره {serial.movement_id} "
                "وصل شده است و برای حفظ سابقه قابل حذف نیست"
            )

        return "این سریال تحویل شده است و برای حفظ سابقه قابل حذف نیست"

    product_id = serial.product_id

    db.session.delete(serial)
    db.session.commit()

    return redirect(f"/serials/{product_id}")



@bp.route("/api/products")
def api_products():

    if "user" not in session:
        return jsonify([])

    q = request.args.get("q", "").strip()

    query = Product.query.filter_by(is_active=1)

    if q:
        query = query.filter(
            Product.name.ilike(f"%{q}%")
        )

    products = query.order_by(Product.name).limit(20).all()

    result = []

    for p in products:

        result.append({
            "id": p.id,
            "name": p.name,
            "qty": p.qty,
            "unit": p.unit or "عدد",
            "status": p.item_status or "نو"
        })

    return jsonify(result)

@bp.route("/api/serials/check")
def api_check_serial():

    if "user" not in session:
        return jsonify({
            "success": False,
            "message": "Unauthorized"
        }), 401

    serial = request.args.get("serial", "").strip()

    if not serial:
        return jsonify({
            "success": False,
            "exists": False
        })

    exists = ProductSerial.query.filter_by(
        serial_number=serial
    ).first()

    return jsonify({
        "success": True,
        "exists": exists is not None
    })

@bp.route("/serials/bulk-add", methods=["POST"])
def bulk_add_serials():

    if "user" not in session:
        return "Unauthorized", 401

    product_id = request.form.get("product_id", type=int)
    serials = request.form.get("serials", "").splitlines()

    added = 0
    duplicated = 0

    for serial in serials:

        serial = serial.strip()

        if not serial:
            continue

        exists = ProductSerial.query.filter_by(
            serial_number=serial
        ).first()

        if exists:
            duplicated += 1
            continue

        db.session.add(
            ProductSerial(
                product_id=product_id,
                serial_number=serial,
                created_by=session["user"]
            )
        )

        added += 1

    db.session.commit()

    if duplicated == 0:

        flash(
            f"{added} سریال با موفقیت ثبت شد.",
            "success"
        )

    elif added == 0:

        flash(
            f"هیچ سریالی ثبت نشد. {duplicated} سریال تکراری بود.",
            "warning"
        )

    else:

        flash(
            f"{added} سریال ثبت شد و {duplicated} سریال تکراری بود.",
            "warning"
        )

    return redirect(f"/serials/{product_id}")

@bp.route("/requests")
def requests_page():

    if "user" not in session:
        return redirect("/")

    status = request.args.get("status", "PENDING")

    query = MaterialRequest.query

    if status == "APPROVED":
        query = query.filter_by(status="APPROVED")

    else:
        query = query.filter_by(status="PENDING")

    requests = (
        query
        .order_by(MaterialRequest.id.desc())
        .all()
    )

    return render_template(
        "requests.html",
        requests=requests,
        status=status,
        projects=Project.query.filter_by(
            is_active=1
        ).order_by(
            Project.name
        ).all(),
        products=Product.query.filter_by(
            is_active=1
        ).order_by(
            Product.name
        ).all()
    )

@bp.route("/requests/new")
def new_request():

    if "user" not in session:
        return redirect("/")

    return redirect("/requests#new-request")

@bp.route("/requests/add", methods=["POST"])
def add_request():

    if "user" not in session:
        return redirect("/")

    req = MaterialRequest(

        requester=session["user"],

        project_name=request.form.get("project_name"),

        priority=request.form.get("priority"),

        location_description=request.form.get("location_description"),

        status="PENDING"

    )

    db.session.add(req)
    db.session.commit()

    product_ids = request.form.getlist("product_id[]")
    qtys = request.form.getlist("qty[]")

    for product_id, qty in zip(product_ids, qtys):

        if not product_id:
            continue

        product = Product.query.get(int(product_id))

        requested_qty = int(qty)

        if product.qty < requested_qty:

            flash(
                f"موجودی کالای «{product.name}» کافی نیست.",
                "danger"
            )

            db.session.delete(req)
            db.session.commit()

            return redirect("/requests")

        db.session.add(

            MaterialRequestItem(

                request_id=req.id,

                product_id=int(product_id),

                qty=requested_qty

            )

        )

    db.session.commit()

    flash("درخواست با موفقیت ثبت شد.", "success")

    return redirect("/requests")

@bp.route("/requests/delete/<int:id>")
def delete_request(id):

    if "user" not in session:
        return redirect("/")

    req = MaterialRequest.query.get_or_404(id)

    if req.status != "PENDING":
        flash(
            "امکان حذف درخواستی که تحویل شده وجود ندارد.",
            "danger"
        )
        return redirect("/requests")

    MaterialRequestItem.query.filter_by(
        request_id=id
    ).delete()

    db.session.delete(req)
    db.session.commit()

    flash(
        "درخواست با موفقیت حذف شد.",
        "success"
    )

    return redirect("/requests")

@bp.route("/requests/view/<int:id>")
def view_request(id):

    if "user" not in session:
        return redirect("/")

    req = MaterialRequest.query.get_or_404(id)

    items = MaterialRequestItem.query.filter_by(
        request_id=id
    ).all()

    products = {
        p.id: p
        for p in Product.query.all()
    }

    return render_template(
        "request_view.html",
        req=req,
        items=items,
        products=products
    )

@bp.route("/requests/edit/<int:id>")
def edit_request(id):

    if "user" not in session:
        return redirect("/")

    req = MaterialRequest.query.get_or_404(id)

    if req.status != "PENDING":
        flash(
            "درخواست تحویل شده قابل ویرایش نیست.",
            "danger"
        )
        return redirect(f"/requests/view/{id}")

    items = MaterialRequestItem.query.filter_by(
        request_id=id
    ).all()

    return render_template(
        "request_edit.html",
        req=req,
        items=items,
        projects=Project.query.filter_by(
            is_active=1
        ).order_by(
            Project.name
        ).all(),
        products=Product.query.filter_by(
            is_active=1
        ).order_by(
            Product.name
        ).all()
    )

@bp.route("/requests/update/<int:id>", methods=["POST"])
def update_request(id):

    if "user" not in session:
        return redirect("/")

    req = MaterialRequest.query.get_or_404(id)

    if req.status != "PENDING":
        flash(
            "درخواست تحویل شده قابل ویرایش نیست.",
            "danger"
        )
        return redirect(f"/requests/view/{id}")

    req.project_name = request.form.get("project_name")

    req.priority = request.form.get("priority")

    req.location_description = request.form.get(
        "location_description"
    )

    item_ids = request.form.getlist("item_id[]")
    qty_list = request.form.getlist("qty[]")

    for item_id, qty in zip(item_ids, qty_list):

        item = MaterialRequestItem.query.get(
            int(item_id)
        )

        if item:

            item.qty = int(qty)

    delete_items = request.form.getlist("delete_item[]")

    if delete_items:

        for item_id in delete_items:

            item = MaterialRequestItem.query.get(
                int(item_id)
            )

            if item:

                count = MaterialRequestItem.query.filter_by(
                    request_id=req.id
                ).count()

            if count > 1:

                db.session.delete(item)

            else:

                flash(
                    "درخواست باید حداقل یک کالا داشته باشد.",
                    "warning"
                )
      
    new_product_id = request.form.get("new_product_id")
    new_qty = request.form.get("new_qty")

    if new_product_id:

        new_product_id = int(new_product_id)
        new_qty = int(new_qty)

        existing = MaterialRequestItem.query.filter_by(
            request_id=req.id,
            product_id=new_product_id
        ).first()

        if existing:

            existing.qty += new_qty

        else:

            db.session.add(
                MaterialRequestItem(
                    request_id=req.id,
                    product_id=new_product_id,
                    qty=new_qty
                )
            )

    db.session.commit()

    flash(
        "تغییرات با موفقیت ذخیره شد.",
        "success"
    )

    return redirect(f"/requests/view/{id}")

@bp.route("/requests/issue/<int:id>")
def issue_request(id):

    if "user" not in session:
        return redirect("/")

    req = MaterialRequest.query.get_or_404(id)

    items = MaterialRequestItem.query.filter_by(
        request_id=id
    ).all()

    products = {
        p.id: p
        for p in Product.query.all()
    }

    available_serials = {}

    for product in products.values():

        if product.has_serial:

            available_serials[product.id] = ProductSerial.query.filter_by(
                product_id=product.id,
                status="IN_STOCK"
            ).order_by(
                ProductSerial.serial_number.asc()
            ).all()

    return render_template(
        "request_issue.html",
        req=req,
        items=items,
        products=products,
        available_serials=available_serials
    )

@bp.route("/product/<int:id>")
def product_view(id):

    if "user" not in session:
        return redirect("/")

    product = Product.query.get_or_404(id)

    movements = (
        Movement.query
        .filter_by(product_id=id)
        .order_by(Movement.timestamp.desc())
        .all()
    )

    serials = (
        ProductSerial.query
        .filter_by(product_id=id)
        .all()
    )

    stock_in = sum(
        m.qty for m in movements
        if m.type == "IN"
    )

    stock_out = sum(
        m.qty for m in movements
        if m.type == "OUT"
    )

    serial_in_stock = sum(
        1 for s in serials
        if s.status == "IN_STOCK"
    )

    return render_template(
        "product_view.html",
        product=product,
        movements=movements,
        serials=serials,
        stock_in=stock_in,
        stock_out=stock_out,
        serial_in_stock=serial_in_stock
    )

@bp.route("/requests/issue/<int:id>", methods=["POST"])
def issue_request_submit(id):

    if "user" not in session:
        return redirect("/")

    req = MaterialRequest.query.get_or_404(id)

    # اگر قبلاً تحویل شده باشد
    if req.status == "APPROVED":
        flash(
            "این درخواست قبلاً تحویل شده است.",
            "warning"
        )
        return redirect(f"/requests/view/{id}")

    items = MaterialRequestItem.query.filter_by(
        request_id=id
    ).all()

    # بررسی موجودی
    for item in items:

        product = Product.query.get(item.product_id)

        if product.qty < item.qty:
            flash(
                f"موجودی کالای '{product.name}' کافی نیست.",
                "danger"
            )
            return redirect(f"/requests/issue/{id}")

    # ثبت تحویل
    for item in items:

        product = Product.query.get(item.product_id)

        # کسر موجودی
        product.qty -= item.qty

        # ثبت تعداد تحویل شده
        item.delivered_qty = item.qty

        # ثبت گردش کالا
        movement = Movement(
            product_id=product.id,
            type="OUT",
            qty=item.qty,
            receiver_name=req.requester,
            project_name=req.project_name,
            description=f"تحویل از درخواست شماره {req.id}",
            created_by=session["user"]
        )

        db.session.add(movement)
        db.session.flush()

        # مدیریت شماره سریال
        if product.has_serial:

            serial_ids = request.form.getlist(
                f"serial_ids_{product.id}[]"
            )

            if len(serial_ids) != item.qty:
                flash(
                    f"برای کالای «{product.name}» باید دقیقاً {item.qty} شماره سریال انتخاب شود.",
                    "danger"
                )
                return redirect(f"/requests/issue/{id}")

            for serial_id in serial_ids:

                serial = ProductSerial.query.get(int(serial_id))

                if (
                    not serial
                    or serial.product_id != product.id
                    or serial.status != "IN_STOCK"
                ):
                    flash(
                        "یکی از شماره سریال‌های انتخاب شده معتبر نیست.",
                        "danger"
                    )
                    return redirect(f"/requests/issue/{id}")

                serial.status = "OUT"
                serial.movement_id = movement.id
                serial.customer_name = req.requester

    # تغییر وضعیت درخواست
    req.status = "APPROVED"
    req.approved_by = session["user"]
    req.approved_at = datetime.utcnow()

    db.session.commit()

    flash(
        "درخواست با موفقیت تحویل شد.",
        "success"
    )

    return redirect(f"/requests/view/{id}")
@bp.route("/projects")
def projects():

    if "user" not in session:
        return redirect("/")

    projects = (
        Project.query
        .order_by(Project.name)
        .all()
    )

    return render_template(
        "projects.html",
        projects=projects
    )

@bp.route("/projects/add", methods=["POST"])
def add_project():

    if "user" not in session:
        return redirect("/")

    if session.get("role") != "admin":
        return "Access Denied"

    name = request.form.get("name", "").strip()
    customer = request.form.get("customer", "").strip()

    if not name:
        flash("نام پروژه وارد نشده است.", "danger")
        return redirect("/projects")

    exists = Project.query.filter_by(name=name).first()

    if exists:
        flash("این پروژه قبلاً ثبت شده است.", "warning")
        return redirect("/projects")

    db.session.add(
        Project(
            name=name,
            customer=customer
        )
    )

    db.session.commit()

    flash("پروژه با موفقیت ثبت شد.", "success")

    return redirect("/projects")
